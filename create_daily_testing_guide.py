from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = "Daily_Software_Testing_Guide.docx"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F7F7F7"
GRID = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRID)
        borders.append(node)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_cell_items(cell, items):
    cell.text = ""
    for i, item in enumerate(items):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.style = "Body Text"
        p.paragraph_format.left_indent = Cm(0.30)
        p.paragraph_format.first_line_indent = Cm(-0.22)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run("• " + item)
        r.font.size = Pt(8.5)


def add_testing_table(doc, title, manual, automation):
    p = doc.add_paragraph()
    p.style = "Heading 1"
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(title)

    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(8.25)
    table.columns[1].width = Cm(8.25)
    set_borders(table)
    set_repeat_header(table.rows[0])
    for i, heading in enumerate(("Manual daily checks", "Automated daily checks")):
        cell = table.rows[0].cells[i]
        cell.width = Cm(8.25)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, NAVY)
        set_cell_margin(cell, 75, 110, 75, 110)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(heading)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for cell in table.rows[1].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margin(cell)
    shade(table.rows[1].cells[0], "FFFFFF")
    shade(table.rows[1].cells[1], PALE_BLUE)
    add_cell_items(table.rows[1].cells[0], manual)
    add_cell_items(table.rows[1].cells[1], automation)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.25)
section.bottom_margin = Cm(1.15)
section.left_margin = Cm(1.6)
section.right_margin = Cm(1.6)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9)
styles["Body Text"].font.name = "Aptos"
styles["Body Text"].font.size = Pt(9)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(21)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.size = Pt(11.5)
styles["Heading 1"].font.bold = True
styles["Heading 1"].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(3)
title.add_run("Daily Frontend and Backend Testing Guide")
remove_paragraph_border(title)

# Word's built-in Title style may carry a theme border. Remove it at style level
# so the heading remains clean in both Word and PDF exports.
title_style_ppr = styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

intro = doc.add_paragraph(style="Body Text")
intro.paragraph_format.space_after = Pt(5)
intro.paragraph_format.line_spacing = 1.05
intro.add_run(
    "This guide defines practical checks that can be performed every day. Manual testing confirms the user experience and unexpected behavior, while automation provides fast, repeatable coverage for stable workflows and APIs."
)

add_testing_table(
    doc,
    "Frontend testing",
    [
        "Verify login, navigation, forms, search, filters, uploads and core user journeys.",
        "Check layout, text, buttons, error messages and responsiveness on supported screen sizes.",
        "Test browser compatibility, keyboard access, focus order and basic accessibility.",
        "Inspect slow loading, broken links, console errors and incorrect API data shown in the interface.",
    ],
    [
        "Run smoke tests for critical journeys after every deployment.",
        "Run UI regression tests for stable screens, validations and role-based access.",
        "Execute cross-browser tests and selected visual comparison checks.",
        "Track page-load timing, failed network calls, JavaScript errors and flaky tests.",
    ],
)

add_testing_table(
    doc,
    "Backend testing",
    [
        "Send valid, invalid, missing and boundary-value requests to important APIs.",
        "Verify status codes, response fields, business rules, authentication and permissions.",
        "Check database changes, duplicate requests, error handling and audit records where access is available.",
        "Review logs for exceptions, timeouts, failed jobs and incorrect device or user mapping.",
    ],
    [
        "Run API smoke and regression suites for critical endpoints.",
        "Validate response schemas, required fields, data types and contract compatibility.",
        "Run controlled load checks for response time, throughput, errors and recovery.",
        "Monitor service health, queues, database connections and scheduled-job failures.",
    ],
)

p = doc.add_paragraph(style="Heading 1")
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(3)
p.add_run("Recommended daily execution")

table = doc.add_table(rows=2, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(4.1), Cm(4.1), Cm(4.1), Cm(4.2)]
set_borders(table)
set_repeat_header(table.rows[0])
headers = ["Start of day", "After a build", "During the day", "End of day"]
values = [
    "Review deployments, open defects and environment health.",
    "Run automated smoke tests, then manually verify critical flows.",
    "Retest fixes and investigate failures with logs and request references.",
    "Publish pass, fail and blocked results with evidence and owners.",
]
for i in range(4):
    table.columns[i].width = widths[i]
    for row in table.rows:
        row.cells[i].width = widths[i]
    head = table.rows[0].cells[i]
    shade(head, NAVY)
    set_cell_margin(head, 70, 80, 70, 80)
    head.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    hp = head.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run(headers[i])
    hr.bold = True
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = RGBColor(255, 255, 255)
    body = table.rows[1].cells[i]
    shade(body, "FFFFFF" if i % 2 == 0 else PALE_GRAY)
    set_cell_margin(body, 85, 90, 85, 90)
    body.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    bp = body.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bp.paragraph_format.space_after = Pt(0)
    br = bp.add_run(values[i])
    br.font.size = Pt(8)

note = doc.add_paragraph(style="Body Text")
note.paragraph_format.space_before = Pt(5)
note.paragraph_format.space_after = Pt(0)
note.paragraph_format.line_spacing = 1.0
r = note.add_run("Daily evidence to keep  ")
r.bold = True
note.add_run("Build and environment, test date, test data, pass or fail status, screenshots or logs, defect ID, severity, owner and retest result.")

doc.core_properties.title = "Daily Frontend and Backend Testing Guide"
doc.core_properties.subject = "Daily manual and automated software testing activities"
doc.core_properties.author = "QA Team"
doc.save(OUTPUT)
print(OUTPUT)
